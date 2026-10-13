"""
Recompone la figura cualitativa de la introducción (fig:c4_example,
`images/qualitative_native_vs_c4_new_sequences.pdf`) a TAMAÑO DE IMPRESIÓN,
con todo el texto vectorial y >= 8 pt.

La versión original (paper/qualitative_examples/build_figure_new_sequences.py)
dibujaba el texto con PIL sobre un lienzo de 1418 px que el paper imprime a
5.4 in (1 px = 0.275 pt): cabeceras a 6 pt, nombre del tracker a 5.5 pt e
IDs a 3.9 pt, y los IDs se salían de su columna sobre las imágenes. Los
frames de origen (frames_extracted/) ya no existen, pero el PNG original es
sin pérdida y tiene geometría exacta (paneles de PANEL px en posiciones
fijas), así que los 12 paneles se recortan de ahí tal cual -- mismas
imágenes, mismo recorte -- y solo se rehace la composición:
  - cabeceras de columna y nombre/hueco de cada tracker como texto vectorial;
  - el ID de cada panel como etiqueta sobre el propio panel (antes eran 2
    líneas en la columna izquierda que no cabían).
Las etiquetas que el renderizador del tracker dibujó DENTRO de los frames
("ID6435 c0 0.79", "REID OK") vienen incrustadas en los píxeles y no se
pueden agrandar sin volver a renderizar los videos; la figura no depende de
ellas (las etiquetas vectoriales llevan la misma información).

Usa la misma infraestructura que paper_fig_reid_module_diagram_pptx.py
(Slide, export vía LibreOffice remoto + Carlito).

Uso:
    python3 paper_fig_qualitative_native_vs_c4.py --export
"""
import argparse
import os

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches

from paper_fig_reid_module_diagram_pptx import (
    FS_MIN, FS_TITLE, INK, LW_TILE, MATCH, MUTED, PRINT_W, REJECT, Slide, export,
)

SRC_PNG = "paper/qualitative_examples/qualitative_native_vs_c4_new_sequences.png"
# geometría del PNG original (build_figure_new_sequences.py)
SRC_PANEL, SRC_HEADER_H, SRC_ROWLABEL_W, SRC_GAP = 420, 46, 150, 4

# Datos de cada caso (build_figure_new_sequences.py::CASES).
CASES = [
    dict(name="ByteTrack", gap="64-frame gap", before="ID6409", native="ID6435", c4="ID6409"),
    dict(name="BoT-SORT", gap="98-frame gap", before="ID49", native="ID61", c4="ID49"),
    dict(name="OC-SORT", gap="87-frame gap", before="ID3", native="ID27", c4="ID3"),
    dict(name="DeepOC-SORT", gap="182-frame gap", before="ID13", native="ID55", c4="ID13"),
]
# El script original dibujaba los IDs de la columna izquierda sin recortar y
# en dos filas se desbordaban ENCIMA del panel "before" (no es contenido del
# frame): ByteTrack ("...(native)"/"...(+C4)", ~28 px) y DeepOC-SORT (la "T",
# ~8 px). Se tapan copiando la franja inmediatamente a la derecha, que es la
# misma superficie lisa. (fila, col): (x0, y0, x1, y1) en coords del panel.
LEAK_FIXES = {
    (0, 0): (0, 216, 31, 258),
    (3, 0): (0, 170, 11, 206),
}

COLS = [
    ("Before occlusion", "3C3C3C"),
    ("After: native (fails)", "B01E1E"),
    ("After: +C4 (recovers)", "198228"),
]


def extract_panels(src_png, out_dir):
    """{(fila, columna): ruta} -- los 12 paneles del PNG original, sin tocar."""
    im = Image.open(src_png).convert("RGB")
    paths = {}
    for r in range(len(CASES)):
        for c in range(len(COLS)):
            x0 = SRC_ROWLABEL_W + c * (SRC_PANEL + SRC_GAP)
            y0 = SRC_HEADER_H + r * (SRC_PANEL + SRC_GAP)
            p = os.path.join(out_dir, f"panel_{r}_{c}.png")
            panel = im.crop((x0, y0, x0 + SRC_PANEL, y0 + SRC_PANEL))
            if (r, c) in LEAK_FIXES:
                lx0, ly0, lx1, ly1 = LEAK_FIXES[(r, c)]
                w = lx1 - lx0
                panel.paste(panel.crop((lx1, ly0, lx1 + w, ly1)), (lx0, ly0))
            panel.save(p)
            paths[(r, c)] = os.path.basename(p)
    return paths


def build(panels, base_dir, out_path):
    sl = Slide(base_dir)
    LABEL_W, GAP, HEAD_H = 1.02, 0.04, 0.22
    FS_NAME, FS_GAP = 11, 9  # la columna de trackers es lo que más se lee: por encima de FS_MIN
    P = (PRINT_W - 0.05 - LABEL_W - 2 * GAP) / 3
    X0 = 0.05 + LABEL_W
    sl.prs.slide_width = Inches(PRINT_W + 0.2)
    sl.prs.slide_height = Inches(HEAD_H + GAP + len(CASES) * (P + GAP) + 0.2)

    for c, (label, color) in enumerate(COLS):
        x = X0 + c * (P + GAP)
        sl.box(x, 0.0, P, HEAD_H, color, color, label, size=FS_TITLE, color="FFFFFF", bold=True,
               lw=LW_TILE, shape=MSO_SHAPE.RECTANGLE)

    for r, case in enumerate(CASES):
        y = HEAD_H + GAP + r * (P + GAP)
        sl.text(0.05, y + P / 2 - 0.23, LABEL_W - 0.05, 0.22, case["name"], FS_NAME, color=INK, bold=True,
                anchor=MSO_ANCHOR.MIDDLE)
        sl.text(0.05, y + P / 2 + 0.03, LABEL_W - 0.05, 0.18, case["gap"], FS_GAP, color="444444")
        tags = [(case["before"], INK), (f"{case['native']} (new ID)", REJECT), (f"{case['c4']} (kept)", MATCH)]
        for c, (tag, tag_color) in enumerate(tags):
            x = X0 + c * (P + GAP)
            sl.image(panels[(r, c)], x, y, P, P, "FFFFFF", LW_TILE)
            t = sl.box(x + 0.04, y + P - 0.22, 0.05 + 0.062 * len(tag), 0.18, "FFFFFF", tag_color, tag,
                       size=FS_MIN, color=tag_color, bold=True, lw=LW_TILE, shape=MSO_SHAPE.RECTANGLE)
            t.text_frame.margin_left = t.text_frame.margin_right = 0
    sl.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--src-png", default=SRC_PNG)
    parser.add_argument("--out", default="paper/qualitative_examples/qualitative_native_vs_c4_new_sequences_print.pptx")
    parser.add_argument("--export", action="store_true", help="genera también <out>.pdf y <out>.png")
    parser.add_argument("--export-host", default=os.environ.get("EXPORT_HOST", "localhost"))
    parser.add_argument("--fonts-dir", default="paper/training_methodology/_fonts")
    args = parser.parse_args()

    base_dir = os.path.join(os.path.dirname(os.path.abspath(args.out)),
                            os.path.splitext(os.path.basename(args.out))[0] + "_assets")
    os.makedirs(base_dir, exist_ok=True)
    panels = extract_panels(args.src_png, base_dir)
    build(panels, base_dir, args.out)
    print(f"[OK] {args.out}")
    if args.export:
        export(args.out, args.export_host, args.fonts_dir)


if __name__ == "__main__":
    main()
